"""Tier 4 - Real-World Application Scenarios (Realistic Workloads).

Simulates production multi-store monthly close cycles, multi-product catalog expansions,
pricing shifts across periods, export generation, and comprehensive audit trail validations.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
from fastapi.testclient import TestClient

from app.core.periods import current_period
from tests.e2e.conftest import make_excel_workbook


def test_workload_end_to_end_monthly_close(
    admin_client: TestClient,
    e2e_env: dict[str, Any],
):
    """TC-LOAD-01: Comprehensive end-to-end monthly close scenario:
    3 stores submitting monthly sales, mid-period pricing revisions, admin review
    and lock, summary calculation, Excel workbook download, backup snapshot,
    catalog Excel import, and full audit trail verification.
    """
    period = current_period()

    # -------------------------------------------------------------------------
    # Phase 1: Setup 3 Stores and 4 Core Catalog Products
    # -------------------------------------------------------------------------
    stores_meta = [
        {"name": "Kyiv Flagship", "login": "kyiv_store", "pass": "KyivPass123!"},
        {"name": "Lviv Market", "login": "lviv_store", "pass": "LvivPass123!"},
        {"name": "Dnipro Central", "login": "dnipro_store", "pass": "DniproPass123!"},
    ]
    created_stores = {}
    for sm in stores_meta:
        res = admin_client.post(
            "/api/v1/admin/stores",
            json={"name": sm["name"], "login": sm["login"], "password": sm["pass"]},
        )
        assert res.status_code == 201
        created_stores[sm["login"]] = res.json()

    products_meta = [
        {"name": "Artisan Rye Bread", "price": "2.20"},     # 220 cents
        {"name": "Pasteurized Milk 2.5%", "price": "1.60"}, # 160 cents
        {"name": "Cream Butter 82%", "price": "4.50"},      # 450 cents
        {"name": "Butter Croissant", "price": "1.25"},      # 125 cents
    ]
    created_products = {}
    for pm in products_meta:
        res = admin_client.post("/api/v1/admin/products", json=pm)
        assert res.status_code == 201
        created_products[pm["name"]] = res.json()

    bread_id = created_products["Artisan Rye Bread"]["id"]
    milk_id = created_products["Pasteurized Milk 2.5%"]["id"]
    butter_id = created_products["Cream Butter 82%"]["id"]
    croissant_id = created_products["Butter Croissant"]["id"]

    # -------------------------------------------------------------------------
    # Phase 2: Mid-Period Price Revisions by Administrator
    # -------------------------------------------------------------------------
    # Increase Bread price from 2.20 to 2.50 (250 cents)
    bread_reprice = admin_client.put(
        f"/api/v1/admin/products/{bread_id}/price",
        json={"price": "2.50"},
    )
    assert bread_reprice.status_code == 200
    assert 250 in bread_reprice.json()["month_prices_cents"]
    assert 220 in bread_reprice.json()["month_prices_cents"]

    # Increase Milk price from 1.60 to 1.80 (180 cents)
    milk_reprice = admin_client.put(
        f"/api/v1/admin/products/{milk_id}/price",
        json={"price": "1.80"},
    )
    assert milk_reprice.status_code == 200
    assert 180 in milk_reprice.json()["month_prices_cents"]
    assert 160 in milk_reprice.json()["month_prices_cents"]

    # -------------------------------------------------------------------------
    # Phase 3: Store Reporting Workload
    # -------------------------------------------------------------------------
    def make_store_client(login: str, password: str) -> TestClient:
        c = TestClient(admin_client.app, base_url="http://testserver", headers={"origin": "http://testserver"})
        c.post("/api/v1/auth/login", json={"login": login, "password": password})
        return c

    # Store 1: Kyiv Flagship
    # Sells Bread at old (220c) and new (250c) prices, plus Milk (180c), Butter (450c), Croissant (125c)
    # 5 Bread @ 220c (1100c) + 10 Bread @ 250c (2500c) + 15 Milk @ 180c (2700c) + 4 Butter @ 450c (1800c) + 20 Croissant @ 125c (2500c)
    # Total = 1100 + 2500 + 2700 + 1800 + 2500 = 10600 cents ($106.00)
    c_kyiv = make_store_client("kyiv_store", "KyivPass123!")
    res_kyiv = c_kyiv.put(
        f"/api/v1/store/reports/{period}",
        json={
            "action": "submit",
            "quantities": [
                {"product_id": bread_id, "price_cents": 220, "quantity": 5},
                {"product_id": bread_id, "price_cents": 250, "quantity": 10},
                {"product_id": milk_id, "price_cents": 180, "quantity": 15},
                {"product_id": butter_id, "price_cents": 450, "quantity": 4},
                {"product_id": croissant_id, "price_cents": 125, "quantity": 20},
            ],
        },
    )
    assert res_kyiv.status_code == 200
    assert res_kyiv.json()["total_cents"] == 10600
    assert res_kyiv.json()["status"] == "submitted"

    # Store 2: Lviv Market
    # 12 Bread @ 250c (3000c) + 10 Milk @ 180c (1800c) + 6 Butter @ 450c (2700c) + 16 Croissant @ 125c (2000c)
    # Total = 3000 + 1800 + 2700 + 2000 = 9500 cents ($95.00)
    c_lviv = make_store_client("lviv_store", "LvivPass123!")
    res_lviv = c_lviv.put(
        f"/api/v1/store/reports/{period}",
        json={
            "action": "submit",
            "quantities": [
                {"product_id": bread_id, "price_cents": 250, "quantity": 12},
                {"product_id": milk_id, "price_cents": 180, "quantity": 10},
                {"product_id": butter_id, "price_cents": 450, "quantity": 6},
                {"product_id": croissant_id, "price_cents": 125, "quantity": 16},
            ],
        },
    )
    assert res_lviv.status_code == 200
    assert res_lviv.json()["total_cents"] == 9500
    assert res_lviv.json()["status"] == "submitted"

    # Store 3: Dnipro Central
    # Saves draft first: 8 Bread @ 250c (2000c) + 20 Milk @ 180c (3600c) + 5 Butter @ 450c (2250c) + 10 Croissant @ 125c (1250c)
    # Total = 2000 + 3600 + 2250 + 1250 = 9100 cents ($91.00)
    c_dnipro = make_store_client("dnipro_store", "DniproPass123!")
    res_dnipro_draft = c_dnipro.put(
        f"/api/v1/store/reports/{period}",
        json={
            "action": "draft",
            "quantities": [
                {"product_id": bread_id, "price_cents": 250, "quantity": 8},
                {"product_id": milk_id, "price_cents": 180, "quantity": 20},
                {"product_id": butter_id, "price_cents": 450, "quantity": 5},
                {"product_id": croissant_id, "price_cents": 125, "quantity": 10},
            ],
        },
    )
    assert res_dnipro_draft.status_code == 200
    assert res_dnipro_draft.json()["total_cents"] == 9100
    assert res_dnipro_draft.json()["status"] == "draft"
    dnipro_report_id = res_dnipro_draft.json()["report_id"]

    # Admin reviews Dnipro's draft and explicitly locks it to submitted
    lock_res = admin_client.patch(
        f"/api/v1/admin/reports/{dnipro_report_id}",
        json={"status": "submitted"},
    )
    assert lock_res.status_code == 200
    assert lock_res.json()["status"] == "submitted"

    # -------------------------------------------------------------------------
    # Phase 4: Financial Summary Matrix & Reconciliation
    # -------------------------------------------------------------------------
    summary = admin_client.get(f"/api/v1/admin/summary?period={period}").json()
    # Expected grand total = 10600 (Kyiv) + 9500 (Lviv) + 9100 (Dnipro) = 29200 cents
    assert summary["grand_total_cents"] >= 29200
    assert len(summary["stores"]) >= 3
    assert len(summary["rows"]) >= 4

    # -------------------------------------------------------------------------
    # Phase 5: Export Verification (Excel and Live Backup Snapshot)
    # -------------------------------------------------------------------------
    # Excel workbook download
    excel_res = admin_client.get(f"/admin/summary.xlsx?period={period}")
    assert excel_res.status_code == 200
    assert "spreadsheetml" in excel_res.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(excel_res.content))
    sheet = wb.active
    assert sheet.title == "Звіт"
    # Check store names in JSON summary matrix
    assert any("Kyiv Flagship" in s["name"] for s in summary["stores"])
    assert any("Lviv Market" in s["name"] for s in summary["stores"])
    assert any("Dnipro Central" in s["name"] for s in summary["stores"])

    # Check products and grand total row in Excel summary workbook
    cell_values = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert any("Artisan Rye Bread" in v for v in cell_values)
    assert any("Разом" in v for v in cell_values)

    # Database backup download
    backup_res = admin_client.get("/admin/backup/download")
    assert backup_res.status_code == 200
    assert backup_res.headers["content-type"] in ("application/octet-stream", "application/sql")
    assert len(backup_res.content) > 0

    # -------------------------------------------------------------------------
    # Phase 6: Catalog Import via Excel
    # -------------------------------------------------------------------------
    import_bytes = make_excel_workbook([
        ("Imported Forest Honey", "7.50"),
        ("Imported Herbal Tea", "3.20"),
    ])
    import_res = admin_client.post(
        "/api/v1/admin/products/import",
        files={"file": ("new_catalog.xlsx", import_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert import_res.status_code == 200
    assert "завершено" in import_res.json()["message"]

    # Verify newly imported items are active in catalogue
    catalogue = admin_client.get("/api/v1/admin/products").json()
    cat_names = [p["name"] for p in catalogue]
    assert "Imported Forest Honey" in cat_names
    assert "Imported Herbal Tea" in cat_names

    # -------------------------------------------------------------------------
    # Phase 7: Comprehensive Audit Trail Verification
    # -------------------------------------------------------------------------
    logs = admin_client.get("/api/v1/admin/logs?limit=50").json()
    log_actions = [log.get("action", "") for log in logs]
    assert any("додано магазин" in a for a in log_actions)
    assert any("зміна ціни" in a or "цін" in a or "змінено" in a for a in log_actions)
    assert any("подання" in a or "звіт" in a for a in log_actions)
    assert any("завантажено бекап" in a for a in log_actions)


def test_workload_high_volume_catalogue_stress(
    admin_client: TestClient,
    store_client: TestClient,
    e2e_env: dict[str, Any],
):
    """TC-LOAD-02: High-volume product catalog stress test:
    Creates 20 distinct products, drafts and submits quantities across all rows,
    and validates high-dimension matrix calculation and Excel workbook generation.
    """
    period = current_period()

    # Create 20 distinct products
    products = []
    for i in range(1, 21):
        price_val = f"{(i * 0.75):.2f}"
        p = admin_client.post(
            "/api/v1/admin/products",
            json={"name": f"Stress Product {i:02d}", "price": price_val},
        ).json()
        products.append(p)

    # Store enters quantities for all 20 products
    quantities = []
    expected_total_cents = 0
    for i, p in enumerate(products, start=1):
        qty = i % 5 + 1
        quantities.append({
            "product_id": p["id"],
            "price_cents": p["price_cents"],
            "quantity": qty,
        })
        expected_total_cents += qty * p["price_cents"]

    # Store submits the 20-item report
    submit_res = store_client.put(
        f"/api/v1/store/reports/{period}",
        json={"action": "submit", "quantities": quantities},
    )
    assert submit_res.status_code == 200
    assert submit_res.json()["total_cents"] >= expected_total_cents
    assert submit_res.json()["status"] == "submitted"

    # Admin checks summary matrix
    summary = admin_client.get(f"/api/v1/admin/summary?period={period}").json()
    assert summary["grand_total_cents"] >= expected_total_cents
    assert len(summary["rows"]) >= 20

    # Admin downloads summary Excel
    xlsx_res = admin_client.get(f"/admin/summary.xlsx?period={period}")
    assert xlsx_res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    sheet = wb.active
    assert sheet.max_row >= 20
