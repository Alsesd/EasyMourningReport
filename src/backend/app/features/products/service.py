"""Product catalogue, pricing timeline engine, and Excel import."""

from __future__ import annotations

import io
from collections import defaultdict
from typing import TYPE_CHECKING

from fastapi import HTTPException, Request
from openpyxl import load_workbook
from sqlmodel import Session, delete, select

import app.core.backup as backup_mod
from app.core.formatting import money
from app.core.periods import current_period
from app.core.validation import clean_name, norm_name, parse_period, to_cents
from app.features.audit.service import record_audit
from app.features.products.models import Product, ProductPrice
from app.features.products.schemas import (
    MutationDTO,
    PricePointDTO,
    ProductCreateDTO,
    ProductDTO,
    ProductUpdateDTO,
)

if TYPE_CHECKING:
    from app.features.auth.models import User

BASE_PERIOD = "0000-00"  # "since always": initial price
History = list[tuple[str, int]]  # [(period, price_cents), ...] sorted by period


def period_month(p: str) -> str:
    """Return YYYY-MM for both canonical and within-month history keys."""
    separators = [i for i in (p.find("_"), p.find("#")) if i >= 0]
    return p[: min(separators)] if separators else p


def _sequence(period_key: str) -> tuple[int, str]:
    """Sequence number and suffix for a price-change key."""
    month = period_month(period_key)
    if len(period_key) == len(month):
        return 0, ""
    suffix = period_key[len(month) + 1 :]
    return (int(suffix), suffix) if suffix.isdigit() else (0, suffix)


def _history_sort_key(period_key: str) -> tuple[str, int, int, str]:
    month = period_month(period_key)
    if len(period_key) == len(month):
        return month, 0, 0, ""
    seq, suffix = _sequence(period_key)
    return month, 1, seq, suffix


def sort_history(history: History) -> History:
    """Sort by month, then by within-month price-change sequence."""
    return sorted(history, key=lambda item: _history_sort_key(item[0]))


def offered_prices(history: History, period: str) -> list[int]:
    """Return every distinct price in force at some point in this month."""
    hist = sort_history(history)
    before = [cents for start, cents in hist if period_month(start) < period]
    current_month = [
        cents for start, cents in hist if period_month(start) == period
    ]
    offered: list[int] = []
    if before:
        offered.append(before[-1])
    for cents in current_month:
        if cents not in offered:
            offered.append(cents)
    return offered


def month_price_timeline(history: History, period: str) -> list[int]:
    """Full monthly price timeline for the admin UI."""
    hist = sort_history(history)
    before = [cents for start, cents in hist if period_month(start) < period]
    current_month = [
        cents for start, cents in hist if period_month(start) == period
    ]
    timeline: list[int] = [before[-1]] if before else []
    timeline.extend(current_month)
    return timeline


def apply_price_change(
    history: History, period: str, new_cents: int
) -> tuple[History, int, int]:
    """Append a price change while preserving earlier changes this month."""
    hist = sort_history(history)
    before = [item for item in hist if period_month(item[0]) < period]
    if not before:
        raise ValueError("history has no price before this period")

    previous_month_price = before[-1][1]
    current_month = [
        (key, cents) for key, cents in hist if period_month(key) == period
    ]
    current_price = (
        current_month[-1][1] if current_month else previous_month_price
    )
    if new_cents == current_price:
        return hist, previous_month_price, current_price

    if not current_month:
        key = period
    else:
        used = {key for key, _ in current_month}
        sequence = (
            max((_sequence(key)[0] for key, _ in current_month), default=0) + 1
        )
        key = f"{period}_{sequence}"
        while key in used:
            sequence += 1
            key = f"{period}_{sequence}"

    new_history = sort_history([*hist, (key, new_cents)])
    return new_history, previous_month_price, new_cents


def latest_price(history: History) -> int:
    ordered = sort_history(history)
    if not ordered:
        raise ValueError("history is empty")
    return ordered[-1][1]


def line_labels(
    derived: list[int], extra: list[int]
) -> list[tuple[int, str]]:
    """(price, label) for each input line of a product."""
    if len(derived) == 2:
        lines = [(derived[0], "old"), (derived[1], "new")]
    elif len(derived) > 2:
        lines = [(derived[0], "old")] + [(p, "new") for p in derived[1:]]
    else:
        lines = [(p, "") for p in derived]
    lines += [(p, "saved") for p in extra if p not in derived]
    return lines


# ---------- Data access & Catalogue Management ----------


def load_history(db: Session) -> dict[int, History]:
    """product id -> [(period, price_cents), ...] sorted by period."""
    out: dict[int, History] = defaultdict(list)
    q = select(
        ProductPrice.product_id, ProductPrice.period, ProductPrice.price_cents
    ).order_by(ProductPrice.period)
    for pid, period, cents in db.execute(q):
        out[pid].append((period, cents))
    return {pid: sort_history(history) for pid, history in out.items()}


def history_of(history: dict[int, History], p: Product) -> History:
    return sort_history(history.get(p.id) or [(BASE_PERIOD, p.price_cents)])


def in_form(p: Product) -> bool:
    """Products that appear in new report forms."""
    return p.active and not p.archived


def save_history(db: Session, pid: int, new: History) -> None:
    rows = {
        r.period: r
        for r in db.scalars(
            select(ProductPrice).where(ProductPrice.product_id == pid)
        )
    }
    wanted = dict(new)
    for period, row in rows.items():
        if period in wanted:
            row.price_cents = wanted[period]
        else:
            db.delete(row)
    for period, cents in wanted.items():
        if period not in rows:
            db.add(
                ProductPrice(product_id=pid, period=period, price_cents=cents)
            )


# Interface contracts for Reports slice


def load_active_products(db: Session) -> list[Product]:
    """Used by reports line builder to load active products."""
    return list(
        db.scalars(
            select(Product).where(Product.active, Product.archived.is_(False))
        ).all()
    )


def load_product_history(db: Session) -> dict[int, list[tuple[str, int]]]:
    """Used by reports summary builder."""
    return load_history(db)


def _get_product_usage(db: Session) -> dict[int, int]:
    # Query reports usage via reports service to respect boundary
    from app.features.reports.service import get_product_report_counts

    return get_product_report_counts(db)


def product_to_dto(
    db: Session,
    product: Product,
    used_count: int = 0,
    period: str | None = None,
    all_history: dict[int, list[tuple[str, int]]] | None = None,
) -> ProductDTO:
    active_period = parse_period(period or current_period())
    history = all_history if all_history is not None else load_history(db)
    product_history = history_of(history, product)
    return ProductDTO(
        id=product.id,
        name=product.name,
        price_cents=product.price_cents,
        active=product.active,
        archived=product.archived,
        used_count=used_count,
        month_prices_cents=month_price_timeline(
            product_history, active_period
        ),
        price_history=[
            PricePointDTO(period=month, price_cents=cents)
            for month, cents in product_history
        ],
    )


def list_products(
    db: Session, period: str | None = None
) -> list[ProductDTO]:
    active_period = parse_period(period)
    used = _get_product_usage(db)
    history = load_history(db)
    products = db.scalars(select(Product).order_by(Product.name)).all()
    return [
        product_to_dto(
            db, product, used.get(product.id, 0), active_period, history
        )
        for product in products
    ]


def create_product(
    db: Session,
    payload: ProductCreateDTO,
    user: User,
    request: Request,
) -> ProductDTO:
    name = clean_name(payload.name)
    cents = to_cents(payload.price)
    product = Product(name=name, price_cents=cents)
    db.add(product)
    db.flush()
    db.add(
        ProductPrice(
            product_id=product.id, period=BASE_PERIOD, price_cents=cents
        )
    )
    record_audit(
        db,
        request,
        user.login,
        user.role,
        "додано товар",
        f"{name} · {money(cents)}",
    )
    db.commit()
    db.refresh(product)
    return product_to_dto(db, product)


def update_product_price(
    db: Session,
    pid: int,
    payload: dict,
    user: User,
    request: Request,
) -> ProductDTO:
    product = db.get(Product, pid)
    if not product or product.archived:
        raise HTTPException(status_code=404, detail="Товар не знайдено.")
    raw_price = str(payload.get("price") or "").strip()
    if not raw_price:
        raise HTTPException(status_code=400, detail="Вкажіть нову ціну.")
    cents = to_cents(raw_price)
    period = current_period()
    history = history_of(load_history(db), product)
    new_history, previous, now = apply_price_change(history, period, cents)
    if new_history != history:
        save_history(db, pid, new_history)
        product.price_cents = latest_price(new_history)
        record_audit(
            db,
            request,
            user.login,
            user.role,
            "змінено ціну товару",
            f"{product.name} · {money(previous)} → {money(now)} з {period}",
        )
        db.commit()
        db.refresh(product)
    return product_to_dto(db, product, _get_product_usage(db).get(pid, 0))


def update_product(
    db: Session,
    pid: int,
    payload: ProductUpdateDTO,
    user: User,
    request: Request,
) -> ProductDTO:
    product = db.get(Product, pid)
    if not product or product.archived:
        raise HTTPException(status_code=404, detail="Товар не знайдено.")
    name = clean_name(payload.name)
    cents = to_cents(payload.price)
    period = current_period()
    history = history_of(load_history(db), product)
    new_history, previous, now = apply_price_change(history, period, cents)
    if new_history != history:
        save_history(db, pid, new_history)
        product.price_cents = latest_price(new_history)
        record_audit(
            db,
            request,
            user.login,
            user.role,
            "змінено ціну товару",
            f"{product.name} · {money(previous)} → {money(now)} з {period}",
        )
    changed = []
    if product.name != name:
        changed.append(f"назва «{product.name}» → «{name}»")
        product.name = name
    if product.active != payload.active:
        product.active = payload.active
        changed.append("увімкнено" if payload.active else "вимкнено")
    if changed:
        record_audit(
            db,
            request,
            user.login,
            user.role,
            "змінено товар",
            f"{name} · " + "; ".join(changed),
        )
    db.commit()
    db.refresh(product)
    return product_to_dto(db, product, _get_product_usage(db).get(pid, 0))


def archive_or_delete_product(
    db: Session,
    pid: int,
    user: User,
    request: Request,
) -> MutationDTO:
    from app.features.reports.service import count_product_reports

    product = db.get(Product, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Товар не знайдено.")
    used = count_product_reports(db, pid)
    if used:
        product.archived, product.active = True, False
        record_audit(
            db,
            request,
            user.login,
            user.role,
            "архівовано товар",
            f"{product.name} · рядків у звітах: {used}",
        )
        db.commit()
        return MutationDTO(
            id=pid,
            archived=True,
            message=f"«{product.name}» переміщено в архів. Дані звітів збережено.",
        )
    backup_mod.backup_db("before-delete")
    name = product.name
    db.execute(delete(ProductPrice).where(ProductPrice.product_id == pid))
    db.execute(delete(Product).where(Product.id == pid))
    record_audit(db, request, user.login, user.role, "видалено товар", name)
    db.commit()
    return MutationDTO(
        id=pid, deleted=True, message=f"Товар «{name}» остаточно видалено."
    )


def restore_product(
    db: Session,
    pid: int,
    user: User,
    request: Request,
) -> ProductDTO:
    product = db.get(Product, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Товар не знайдено.")
    product.archived, product.active = False, True
    record_audit(
        db, request, user.login, user.role, "відновлено товар", product.name
    )
    db.commit()
    db.refresh(product)
    return product_to_dto(db, product, _get_product_usage(db).get(pid, 0))


def parse_products_excel(
    payload: bytes, db: Session, user: User, request: Request
) -> str:
    """Parse products from Excel, insert new rows, and return a user-facing summary."""
    try:
        workbook = load_workbook(
            io.BytesIO(payload), read_only=True, data_only=True
        )
        sheet = workbook.active
        parsed = []
        errors = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_col=1, max_col=2, values_only=True), start=1
        ):
            raw_name, raw_price = row
            if raw_name is None and raw_price is None:
                continue
            if row_number == 1 and str(raw_name or "").strip().casefold() in {
                "name",
                "назва",
                "товар",
                "product",
            }:
                continue
            try:
                name = clean_name(str(raw_name or ""))
                if raw_price is None or isinstance(raw_price, bool):
                    raise ValueError("не вказана ціна")
                cents = to_cents(str(raw_price))
                parsed.append((row_number, name, cents))
            except Exception as exc:
                errors.append(f"рядок {row_number}: {exc}")
        workbook.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            400,
            "Не вдалося прочитати Excel-файл. Перевірте, що це коректний .xlsx.",
        )
    if errors:
        raise HTTPException(
            400, "Імпорт скасовано — виправте файл:\n" + "\n".join(errors[:20])
        )
    if not parsed:
        raise HTTPException(400, "У файлі немає товарів для імпорту")

    existing_names = {norm_name(p.name) for p in db.scalars(select(Product))}
    seen = set()
    created, skipped = [], 0
    try:
        for row_number, name, cents in parsed:
            key = norm_name(name)
            if key in existing_names or key in seen:
                skipped += 1
                continue
            p = Product(name=name, price_cents=cents)
            db.add(p)
            db.flush()
            db.add(
                ProductPrice(
                    product_id=p.id, period=BASE_PERIOD, price_cents=cents
                )
            )
            created.append((name, cents))
            seen.add(key)
        if created:
            record_audit(
                db,
                request,
                user.login,
                user.role,
                "масово додано товари",
                f"створено: {len(created)}, пропущено дублікатів: {skipped}",
            )
        db.commit()
    except Exception:
        db.rollback()
        raise
    if not created:
        return f"Нових товарів не додано: усі {skipped} рядків уже існують."
    return f"Імпорт завершено: додано {len(created)} товарів, пропущено дублікатів: {skipped}."


__all__ = [
    "BASE_PERIOD",
    "History",
    "apply_price_change",
    "archive_or_delete_product",
    "create_product",
    "history_of",
    "in_form",
    "latest_price",
    "line_labels",
    "list_products",
    "load_active_products",
    "load_history",
    "load_product_history",
    "month_price_timeline",
    "offered_prices",
    "parse_products_excel",
    "period_month",
    "product_to_dto",
    "restore_product",
    "save_history",
    "sort_history",
    "update_product",
    "update_product_price",
]
