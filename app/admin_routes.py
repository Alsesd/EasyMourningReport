"""Admin settings, store/product management, logs, and report routes."""


def register_admin_routes(app, deps):
    import io

    from fastapi import Depends, File, Form, HTTPException, Request, UploadFile
    from fastapi.responses import Response, StreamingResponse
    from openpyxl import Workbook, load_workbook
    from sqlalchemy import delete, func, select
    from sqlalchemy.exc import IntegrityError
    from sqlalchemy.orm import Session

    from .models import AuditLog, Product, ProductPrice, Report, ReportItem, Store, User

    ACTIONS, BASE_PERIOD, admin_period, admin_required, apply_price_change, audit, backup_db, build_lines, build_summary, clean_login, clean_name, current_period, done, FormError, get_db, hash_pw, history_of, in_form, latest_price, load_history, money, new_password, norm_login, norm_name, offered_prices, parse_period, render, report_total, same_name, save_history, to_cents, valid_login = (deps[name] for name in ["ACTIONS","BASE_PERIOD","admin_period","admin_required","apply_price_change","audit","backup_db","build_lines","build_summary","clean_login","clean_name","current_period","done","FormError","get_db","hash_pw","history_of","in_form","latest_price","load_history","money","new_password","norm_login","norm_name","offered_prices","parse_period","render","report_total","same_name","save_history","to_cents","valid_login"])

    # ---------- admin: settings ----------
    @app.get("/admin")
    def admin_page(request: Request, user: User = Depends(admin_required), db: Session = Depends(get_db)):
        period = current_period()
        history = load_history(db)
        used = {pid: n for pid, n in db.execute(
            select(ReportItem.product_id, func.count(ReportItem.id)).group_by(ReportItem.product_id))}
        reports = {sid: n for sid, n in db.execute(select(Report.store_id, func.count(Report.id)).group_by(Report.store_id))}
        products, inactive_products, archived_products = [], [], []
        for p in db.scalars(select(Product).order_by(Product.name)):
            if p.archived:
                archived_products.append({"p": p, "used": used.get(p.id, 0)})
                continue
            lines = offered_prices(history_of(history, p), period)
            item = {"p": p, "used": used.get(p.id, 0), "prev": lines[0],
                    "change": (lines[0], lines[1]) if len(lines) == 2 else None}
            (products if p.active else inactive_products).append(item)
        stores, archived_stores = [], []
        for u in db.scalars(select(User).where(User.role == "store").order_by(User.login)):
            if not u.store:
                continue
            x = {"store": u.store, "user": u, "reports": reports.get(u.store_id, 0)}
            (archived_stores if u.store.archived else stores).append(x)
        stores.sort(key=lambda x: x["store"].name)
        archived_stores.sort(key=lambda x: x["store"].name)
        ctx = {"user": user, "products": products, "inactive_products": inactive_products,
               "archived_products": archived_products,
               "stores": stores, "archived_stores": archived_stores}
        return render(request, "admin.html", ctx)
    
    
    @app.post("/admin/products/import")
    async def product_import(request: Request, file: UploadFile = File(...),
                             user: User = Depends(admin_required), db: Session = Depends(get_db)):
        """Create products from an Excel workbook: column A=name, column B=price."""
        filename = (file.filename or "").lower()
        if not filename.endswith(".xlsx"):
            raise HTTPException(400, "Завантажте файл Excel у форматі .xlsx")
        payload = await file.read()
        if not payload or len(payload) > 10 * 1024 * 1024:
            raise HTTPException(400, "Файл порожній або перевищує 10 МБ")
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            sheet = workbook.active
            parsed = []
            errors = []
            for row_number, row in enumerate(sheet.iter_rows(min_col=1, max_col=2, values_only=True), start=1):
                raw_name, raw_price = row
                if raw_name is None and raw_price is None:
                    continue
                if row_number == 1 and str(raw_name or "").strip().casefold() in {"name", "назва", "товар", "product"}:
                    continue
                try:
                    name = clean_name(str(raw_name or ""))
                    if raw_price is None or isinstance(raw_price, bool):
                        raise ValueError("не вказана ціна")
                    cents = to_cents(str(raw_price))
                    parsed.append((row_number, name, cents))
                except Exception as exc:
                    errors.append(f"рядок {row_number}: {exc}")
            workbook.close()
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(400, "Не вдалося прочитати Excel-файл. Перевірте, що це коректний .xlsx.")
        if errors:
            raise HTTPException(400, "Імпорт скасовано — виправте файл:\n" + "\n".join(errors[:20]))
        if not parsed:
            raise HTTPException(400, "У файлі немає товарів для імпорту")
        existing_names = {norm_name(p.name) for p in db.scalars(select(Product))}
        seen = set()
        created, skipped = [], 0
        try:
            for row_number, name, cents in parsed:
                key = norm_name(name)
                if key in existing_names or key in seen:
                    skipped += 1
                    continue
                p = Product(name=name, price_cents=cents)
                db.add(p)
                db.flush()
                db.add(ProductPrice(product_id=p.id, period=BASE_PERIOD, price_cents=cents))
                created.append((name, cents))
                seen.add(key)
            if created:
                audit(db, request, user.login, user.role, "масово додано товари",
                      f"створено: {len(created)}, пропущено дублікатів: {skipped}")
            db.commit()
        except Exception:
            db.rollback()
            raise
        if not created:
            return done(request, "/admin#products", f"Нових товарів не додано: усі {skipped} рядків уже існують.")
        return done(request, "/admin#products",
                    f"Імпорт завершено: додано {len(created)} товарів, пропущено дублікатів: {skipped}.")

    @app.post("/admin/products")
    def product_add(request: Request, name: str = Form(), price: str = Form(), user: User = Depends(admin_required),
                    db: Session = Depends(get_db)):
        name, cents = clean_name(name), to_cents(price)
        p = Product(name=name, price_cents=cents)
        db.add(p)
        db.flush()
        db.add(ProductPrice(product_id=p.id, period=BASE_PERIOD, price_cents=cents))
        audit(db, request, user.login, user.role, "додано товар", f"{name} · {money(cents)}")
        db.commit()
        return done(request, f"/admin#p{p.id}", f"Товар «{name}» додано з ціною {money(cents)}.")
    
    
    @app.post("/admin/products/{pid}")
    def product_edit(request: Request, pid: int, name: str = Form(), price: str = Form(), active: str | None = Form(None),
                     user: User = Depends(admin_required), db: Session = Depends(get_db)):
        p = db.get(Product, pid)
        if not p or p.archived:
            raise HTTPException(404)
        name, cents, is_active = clean_name(name), to_cents(price), active is not None
        period = current_period()
        said: list[str] = []
    
        history = history_of(load_history(db), p)
        before_now = offered_prices(history, period)[-1]  # price in force right now
        new_history, prev_price, now_price = apply_price_change(history, period, cents)
        if new_history != history:
            save_history(db, pid, new_history)
            p.price_cents = latest_price(new_history)
            if len(offered_prices(new_history, period)) == 2:
                said.append(f"Ціну змінено: {money(before_now)} → {money(now_price)} з {period}. У звітах {period} усі "
                            f"магазини (зокрема нові) бачать два рядки: {money(prev_price)} і {money(now_price)}. "
                            f"У наступних місяцях буде лише {money(now_price)}.")
            else:
                said.append(f"Зміну ціни за {period} скасовано: ціна знову {money(now_price)}, "
                            f"у звітах {period} один рядок.")
            audit(db, request, user.login, user.role, "змінено ціну товару",
                  f"{p.name} · {money(before_now)} → {money(now_price)} з {period}")
    
        notes = []
        if name != p.name:
            notes.append(f"назва «{p.name}» → «{name}»")
            p.name = name
        if is_active != p.active:
            notes.append("увімкнено" if is_active else "вимкнено (не показується в нових звітах, дані збережено)")
            p.active = is_active
        if notes:
            said.append("Збережено: " + "; ".join(notes) + ".")
            audit(db, request, user.login, user.role, "змінено товар", f"{p.name} · " + "; ".join(notes))
        db.commit()
        return done(request, f"/admin#p{pid}", " ".join(said) if said else "Змін немає: нічого не збережено.",
                    "ok" if said else "info")
    
    
    @app.post("/admin/products/{pid}/delete")
    def product_delete(request: Request, pid: int, user: User = Depends(admin_required), db: Session = Depends(get_db)):
        p = db.get(Product, pid)
        if not p:
            raise HTTPException(404)
        used = db.scalar(select(func.count(ReportItem.id)).where(ReportItem.product_id == pid)) or 0
        if used:  # never touch report data: move to the archive instead
            p.archived, p.active = True, False
            audit(db, request, user.login, user.role, "архівовано товар", f"{p.name} · рядків у звітах: {used}")
            db.commit()
            return done(request, "/admin#archive", f"«{p.name}» є у звітах ({used} рядків), тому його не видалено, "
                        "а переміщено в архів. Усі дані звітів збережено; товар можна відновити.")
        backup_db("before-delete")
        name = p.name
        db.execute(delete(ProductPrice).where(ProductPrice.product_id == pid))
        db.execute(delete(Product).where(Product.id == pid))
        audit(db, request, user.login, user.role, "видалено товар", name)
        db.commit()
        return done(request, "/admin#products", f"Товар «{name}» видалено: у звітах його не було.")
    
    
    @app.post("/admin/products/{pid}/restore")
    def product_restore(request: Request, pid: int, user: User = Depends(admin_required), db: Session = Depends(get_db)):
        p = db.get(Product, pid)
        if not p:
            raise HTTPException(404)
        p.archived, p.active = False, True
        audit(db, request, user.login, user.role, "відновлено товар", p.name)
        db.commit()
        return done(request, f"/admin#p{pid}", f"Товар «{p.name}» відновлено й активовано.")
    
    
    def check_unique(db: Session, name: str, login: str, store_id: int | None = None, user_id: int | None = None) -> None:
        """Store names and logins must be unique; case and extra spaces do not count as a difference."""
        for s in db.scalars(select(Store)):
            if s.id != store_id and same_name(s.name, name):
                raise FormError(f"Магазин із назвою «{s.name}» уже існує" + (" (в архіві)" if s.archived else ""))
        other = db.scalar(select(User).where(User.login == login, User.id != (user_id or 0)))
        if other:
            raise FormError(f"Логін «{login}» уже зайнятий. Логіни не залежать від регістру: Store1 і store1 — той самий.")
    
    
    @app.post("/admin/stores")
    def store_add(request: Request, name: str = Form(), login: str = Form(), password: str = Form(),
                  user: User = Depends(admin_required), db: Session = Depends(get_db)):
        name, login = clean_name(name, "Назва магазину"), clean_login(login)
        new_password(password)
        check_unique(db, name, login)
        store = Store(name=name)
        db.add(store)
        try:
            db.flush()
            db.add(User(login=login, password_hash=hash_pw(password), role="store", store_id=store.id))
            audit(db, request, user.login, user.role, "додано магазин", f"{name} · логін {login}")
            db.commit()
        except IntegrityError:
            db.rollback()
            raise FormError("Магазин із такою назвою або логін уже існує")
        return done(request, f"/admin#s{store.id}", f"Магазин «{name}» додано. Логін для входу: {login}.")
    
    
    @app.post("/admin/stores/{sid}/edit")
    def store_edit(request: Request, sid: int, name: str = Form(), login: str = Form(),
                   user: User = Depends(admin_required), db: Session = Depends(get_db)):
        store = db.get(Store, sid)
        store_user = db.scalar(select(User).where(User.store_id == sid, User.role == "store"))
        if not store or not store_user:
            raise HTTPException(404)
        name, login = clean_name(name, "Назва магазину"), clean_login(login)
        check_unique(db, name, login, store_id=sid, user_id=store_user.id)
        notes = []
        if name != store.name:
            notes.append(f"назва «{store.name}» → «{name}»")
            store.name = name
        if login != store_user.login:
            notes.append(f"логін {store_user.login} → {login}")
            store_user.login = login
        if not notes:
            return done(request, f"/admin#s{sid}", "Змін немає: нічого не збережено.", "info")
        try:
            audit(db, request, user.login, user.role, "змінено магазин", "; ".join(notes))
            db.commit()
        except IntegrityError:
            db.rollback()
            raise FormError("Магазин із такою назвою або логін уже існує")
        return done(request, f"/admin#s{sid}", "Збережено: " + "; ".join(notes) + ". Пароль і звіти не змінились.")
    
    
    @app.post("/admin/stores/{sid}/delete")
    def store_delete(request: Request, sid: int, user: User = Depends(admin_required), db: Session = Depends(get_db)):
        s = db.get(Store, sid)
        if not s:
            raise HTTPException(404)
        reports = db.scalar(select(func.count(Report.id)).where(Report.store_id == sid)) or 0
        if reports:  # never touch report data: switch the store off instead
            s.archived = True
            audit(db, request, user.login, user.role, "архівовано магазин", f"{s.name} · звітів: {reports}")
            db.commit()
            return done(request, "/admin#archive", f"«{s.name}» має звіти ({reports}), тому його не видалено, а "
                        "переміщено в архів: вхід вимкнено, усі звіти збережено. Магазин можна відновити.")
        backup_db("before-delete")
        name = s.name
        db.execute(delete(User).where(User.store_id == sid))
        db.execute(delete(Store).where(Store.id == sid))
        audit(db, request, user.login, user.role, "видалено магазин", name)
        db.commit()
        return done(request, "/admin#stores", f"Магазин «{name}» видалено: звітів у нього не було.")
    
    
    @app.post("/admin/stores/{sid}/restore")
    def store_restore(request: Request, sid: int, user: User = Depends(admin_required), db: Session = Depends(get_db)):
        s = db.get(Store, sid)
        if not s:
            raise HTTPException(404)
        s.archived = False
        audit(db, request, user.login, user.role, "відновлено магазин", s.name)
        db.commit()
        return done(request, f"/admin#s{sid}", f"Магазин «{s.name}» відновлено: вхід знову працює.")
    
    
    @app.post("/admin/stores/{sid}/password")
    def store_password(request: Request, sid: int, password: str = Form(), user: User = Depends(admin_required),
                       db: Session = Depends(get_db)):
        store_user = db.scalar(select(User).where(User.store_id == sid, User.role == "store"))
        if not store_user:
            raise HTTPException(404)
        store_user.password_hash = hash_pw(new_password(password))
        audit(db, request, user.login, user.role, "змінено пароль магазину", store_user.login)
        db.commit()
        return done(request, f"/admin#s{sid}", f"Пароль для входу «{store_user.login}» змінено.")
    
    
    @app.post("/admin/password")
    def admin_password(request: Request, password: str = Form(), user: User = Depends(admin_required),
                       db: Session = Depends(get_db)):
        user.password_hash = hash_pw(new_password(password))
        audit(db, request, user.login, user.role, "змінено пароль адміністратора")
        db.commit()
        return done(request, "/admin#password", "Ваш пароль змінено.")
    
    
    # ---------- admin: action log ----------
    @app.get("/admin/logs")
    def logs_page(request: Request, login: str | None = None, action: str | None = None,
                  user: User = Depends(admin_required), db: Session = Depends(get_db)):
        q = select(AuditLog)
        if login:
            q = q.where(AuditLog.login.contains(login.strip().lower()))
        if action:
            q = q.where(AuditLog.action == action)
        logs = db.scalars(q.order_by(AuditLog.id.desc()).limit(200)).all()
        ctx = {"user": user, "logs": logs, "actions": ACTIONS, "login": login or "", "action": action or ""}
        return render(request, "logs.html", ctx)
    
    
    # ---------- admin: individual reports, lock / unlock ----------
    @app.get("/admin/reports")
    def reports_page(request: Request, period: str | None = None, user: User = Depends(admin_required),
                     db: Session = Depends(get_db)):
        period = admin_period(request, period)
        reports = {r.store_id: r for r in db.scalars(select(Report).where(Report.period == period))}
        rows = [{"store": s, "report": reports.get(s.id), "total": report_total(reports[s.id]) if s.id in reports else 0}
                for s in db.scalars(select(Store).order_by(Store.name)) if not s.archived or s.id in reports]
        return render(request, "reports.html", {"user": user, "period": period, "rows": rows})
    
    
    @app.get("/admin/reports/{rid}")
    def report_view(rid: int, request: Request, user: User = Depends(admin_required), db: Session = Depends(get_db)):
        report = db.get(Report, rid)
        if not report:
            raise HTTPException(404)
        request.session["period"] = report.period  # the tabs keep showing this month
        rows, inactive, inactive_total = build_lines(db, report.period, report.items)
        lines = [{"name": r["name"], "label": i["label"], "price": i["price"], "qty": i["qty"],
                  "sum": i["qty"] * i["price"]}
                 for r in rows for i in r["inputs"] if i["qty"] or r["repriced"]]
        ctx = {"user": user, "report": report, "period": report.period, "lines": lines, "inactive_rows": inactive,
               "total": report_total(report)}
        return render(request, "report_view.html", ctx)
    
    
    @app.post("/admin/reports/{rid}/status")
    def report_status(request: Request, rid: int, status: str = Form(), user: User = Depends(admin_required),
                      db: Session = Depends(get_db)):
        report = db.get(Report, rid)
        if not report or status not in ("draft", "submitted"):
            raise HTTPException(400, "Некоректний запит")
        report.status = status
        audit(db, request, user.login, user.role,
              "заблоковано звіт" if status == "submitted" else "розблоковано звіт",
              f"{report.store.name if report.store else '?'} · {report.period}")
        db.commit()
        return done(request, f"/admin/reports/{rid}",
                    "Звіт заблоковано: магазин більше не може його змінити." if status == "submitted"
                    else "Звіт розблоковано: магазин може його редагувати й надіслати знову.")
    
    
    # ---------- admin: monthly summary ----------
    @app.get("/admin/summary")
    def summary_page(request: Request, period: str | None = None, submitted_only: bool = False,
                     user: User = Depends(admin_required), db: Session = Depends(get_db)):
        period = admin_period(request, period)
        stores, rows, store_sums, grand = build_summary(db, period, submitted_only)
        ctx = {"user": user, "period": period, "stores": stores, "rows": rows, "store_sums": store_sums, "grand": grand,
               "submitted_only": submitted_only}
        return render(request, "summary.html", ctx)
    
    
    @app.get("/admin/summary.xlsx")
    def summary_xlsx(period: str | None = None, submitted_only: bool = False, user: User = Depends(admin_required),
                     db: Session = Depends(get_db)):
        period = parse_period(period)
        stores, rows, store_sums, grand = build_summary(db, period, submitted_only)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Звіт"
        sheet.append(["Назва", "Ціна", "Звітована кількість", "Сума"])
        for row in rows:
            sheet.append([row["name"], row["price"] / 100, row["qty"], row["sum"] / 100])
        sheet.append(["Разом", None, sum(row["qty"] for row in rows), grand / 100])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:D{sheet.max_row - 1}" if sheet.max_row > 1 else "A1:D1"
        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 16
        sheet.column_dimensions["C"].width = 24
        sheet.column_dimensions["D"].width = 18
        for cell in sheet["B"][1:] + sheet["D"][1:]:
            cell.number_format = '#,##0.00'
        for cell in sheet[1]:
            cell.font = cell.font.copy(bold=True)
        for cell in sheet[sheet.max_row]:
            cell.font = cell.font.copy(bold=True)
        out = io.BytesIO()
        workbook.save(out)
        out.seek(0)
        return StreamingResponse(
            out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="zvit-{period}.xlsx"'},
        )
